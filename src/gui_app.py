#!/usr/bin/env python3
"""
GUI Application Module

This module provides a graphical user interface for the OCR application using tkinter
with drag-and-drop functionality via tkinterdnd2. It includes a main window with file
drop zone, results display area, save functionality, progress indicators, and
advanced features like batch queue processing, language selection, confidence slider,
preprocessing preview, engine selector, and dark theme.
"""

import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext
import logging
from pathlib import Path
from typing import List, Optional, Dict, Any
import threading
import queue
import os
import io

try:
    from tkinterdnd2 import TkinterDnD, DND_FILES
    TKINTERDND_AVAILABLE = True
except ImportError:
    TKINTERDND_AVAILABLE = False
    logging.error("tkinterdnd2 not available. Drag-and-drop functionality will be disabled. Install with: pip install tkinterdnd2")

from .ocr_engine import OCREngine, get_available_engines
from .pdf_handler import PDFHandler
from .config_loader import get_config_parser, apply_cli_overrides
from .batch_queue import BatchQueueManager, TaskStatus
from .exceptions import (
    OCRError, OCREngineNotAvailableError, OCRProcessingError,
    ImageProcessingError, PDFProcessingError, MemoryError, OCREngineFallbackError
)

# Configure logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# Available languages for the language selector
AVAILABLE_LANGUAGES = [
    ('en', 'English'),
    ('fr', 'French'),
    ('de', 'German'),
    ('es', 'Spanish'),
    ('it', 'Italian'),
    ('pt', 'Portuguese'),
    ('nl', 'Dutch'),
    ('ru', 'Russian'),
    ('ja', 'Japanese'),
    ('ko', 'Korean'),
    ('zh', 'Chinese'),
    ('ar', 'Arabic'),
    ('hi', 'Hindi'),
    ('tr', 'Turkish'),
    ('pl', 'Polish'),
    ('cs', 'Czech'),
    ('sv', 'Swedish'),
    ('da', 'Danish'),
    ('fi', 'Finnish'),
    ('el', 'Greek'),
]


def _format_gui_tables(tables: list) -> str:
    """Format extracted tables as text for display in the GUI."""
    if not tables:
        return ""

    output_parts = []
    for i, table in enumerate(tables):
        try:
            text = table.df.to_string(index=False)
            output_parts.append(f"--- Table {i + 1} ---\n{text}")
        except Exception as e:
            logger.warning(f"Failed to format table {i + 1}: {e}")
    return '\n\n'.join(output_parts)


class OCRGUIApp:
    """
    Main GUI application class for OCR processing.

    Features:
    - Drag-and-drop file support
    - Progress indication during processing
    - Results display with save functionality
    - Error handling and status updates
    - Support for images and PDFs
    - Language selector dropdown
    - Confidence threshold slider
    - Preprocessing preview canvas
    - Engine selector
    - Dark theme support
    - Batch queue processing
    - Memory monitor
    """

    def __init__(self, root: tk.Tk):
        """
        Initialize the GUI application.

        Args:
            root: Tkinter root window
        """
        self.root = root
        self.root.title("OCR Text Extractor")
        self.root.geometry("1000x800")
        self.root.minsize(800, 600)

        # Apply dark theme
        self.style = ttk.Style()
        self._apply_dark_theme()

        # Load configuration
        self.config = get_config_parser()

        # Initialize OCR components with Phase 2 features
        self.ocr_engine = OCREngine(
            languages=self.config.get('languages', ['en']),
            confidence_threshold=self.config.get('confidence_threshold', 0.5),
            use_gpu=self.config.get('use_gpu', False),
            preprocess_images=self.config.get('preprocess_images', True),
            memory_limit_mb=self.config.get('memory_limit_mb'),
            deskew=self.config.get('deskew', False),
            clahe=self.config.get('clahe', False),
            upscale=self.config.get('upscale', False),
            remove_bg=self.config.get('remove_bg', False),
            preserve_layout=self.config.get('preserve_layout', False),
            auto_select=self.config.get('auto_select', True),
        )
        self.pdf_handler = PDFHandler(ocr_engine=self.ocr_engine)

        # Processing queue for thread communication
        self.processing_queue = queue.Queue()

        # Current processing state
        self.current_file = None
        self.processing = False
        self.preview_image_original = None
        self.preview_image_processed = None

        # Log messages list
        self.log_messages: List[str] = []

        # Batch queue manager
        self.batch_manager: Optional[BatchQueueManager] = None

        # Setup UI
        self.setup_ui()

        # Setup drag and drop if available
        if TKINTERDND_AVAILABLE:
            self.setup_drag_drop()

        # Start queue processing and memory monitor
        self.root.after(100, self.process_queue)
        self.root.after(1000, self.update_memory_usage)

    def _apply_dark_theme(self):
        """Apply dark theme using ttk.Style with clam theme."""
        try:
            self.style.theme_use('clam')
        except tk.TclError:
            pass

        # Dark theme colors
        bg_color = '#2b2b2b'
        fg_color = '#ffffff'
        select_color = '#404040'
        button_color = '#3c3c3c'
        entry_bg = '#3c3c3c'
        text_bg = '#1e1e1e'
        text_fg = '#d4d4d4'

        self.style.configure('TFrame', background=bg_color)
        self.style.configure('TLabel', background=bg_color, foreground=fg_color)
        self.style.configure('TButton', background=button_color, foreground=fg_color)
        self.style.configure('TCheckbutton', background=bg_color, foreground=fg_color)
        self.style.configure('TRadiobutton', background=bg_color, foreground=fg_color)
        self.style.configure('TLabelframe', background=bg_color, foreground=fg_color)
        self.style.configure('TLabelframe.Label', background=bg_color, foreground=fg_color)
        self.style.configure('TNotebook', background=bg_color)
        self.style.configure('TNotebook.Tab', background=button_color, foreground=fg_color)
        self.style.configure('TProgressbar', background=select_color, troughcolor=bg_color)
        self.style.configure('Horizontal.TScale', background=bg_color)
        self.style.configure('TCombobox', fieldbackground=entry_bg, foreground=fg_color, background=button_color)
        self.style.configure('Treeview', background=text_bg, foreground=text_fg, fieldbackground=text_bg)
        self.style.configure('Treeview.Heading', background=button_color, foreground=fg_color)

        self.root.configure(bg=bg_color)

    def _log(self, message: str):
        """Add a log message to the log list and queue."""
        import datetime
        timestamp = datetime.datetime.now().strftime('%H:%M:%S')
        entry = f"[{timestamp}] {message}"
        self.log_messages.append(entry)
        self.processing_queue.put(('log', entry))

    def update_memory_usage(self):
        """Update memory usage label in the status bar."""
        try:
            import psutil
            process = psutil.Process(os.getpid())
            mem_mb = process.memory_info().rss / (1024 * 1024)
            mem_percent = process.memory_info().percent
            if hasattr(self, 'memory_label'):
                self.memory_label.config(text=f"Memory: {mem_mb:.0f} MB ({mem_percent:.1f}%)")
        except Exception:
            if hasattr(self, 'memory_label'):
                self.memory_label.config(text="Memory: N/A")
        self.root.after(1000, self.update_memory_usage)

    def setup_ui(self):
        """Setup the main user interface."""
        # Create main frame
        main_frame = ttk.Frame(self.root, padding="10")
        main_frame.pack(fill=tk.BOTH, expand=True)

        # Title
        title_label = ttk.Label(main_frame, text="OCR Text Extractor",
                               font=("Arial", 18, "bold"))
        title_label.pack(pady=(0, 10))

        # Top controls frame
        top_controls = ttk.Frame(main_frame)
        top_controls.pack(fill=tk.X, pady=(0, 10))

        # Language selector
        lang_frame = ttk.LabelFrame(top_controls, text="Language", padding="5")
        lang_frame.pack(side=tk.LEFT, padx=(0, 10))

        self.language_var = tk.StringVar(value=self.config.get('languages', ['en'])[0])
        self.language_combo = ttk.Combobox(lang_frame, textvariable=self.language_var,
                                           values=[f"{code} - {name}" for code, name in AVAILABLE_LANGUAGES],
                                           state='readonly', width=20)
        self.language_combo.pack()
        self.language_combo.bind('<<ComboboxSelected>>', self._on_language_change)

        # Confidence slider
        conf_frame = ttk.LabelFrame(top_controls, text="Confidence Threshold", padding="5")
        conf_frame.pack(side=tk.LEFT, padx=(0, 10))

        self.confidence_var = tk.DoubleVar(value=self.config.get('confidence_threshold', 0.5))
        self.confidence_slider = ttk.Scale(conf_frame, from_=0.0, to=1.0,
                                           variable=self.confidence_var,
                                           orient=tk.HORIZONTAL, length=150,
                                           command=self._on_confidence_change)
        self.confidence_slider.pack(side=tk.LEFT)
        self.confidence_label = ttk.Label(conf_frame, text=f"{self.confidence_var.get():.2f}")
        self.confidence_label.pack(side=tk.LEFT, padx=(5, 0))

        # Engine selector
        engine_frame = ttk.LabelFrame(top_controls, text="Engine", padding="5")
        engine_frame.pack(side=tk.LEFT, padx=(0, 10))

        self.engine_var = tk.StringVar(value='auto')
        engines = get_available_engines()
        engine_options = [('auto', 'Auto')] + [(e, e.capitalize()) for e in engines]
        for val, text in engine_options:
            ttk.Radiobutton(engine_frame, text=text, variable=self.engine_var,
                           value=val).pack(side=tk.LEFT, padx=2)

        # Drop zone frame
        drop_frame = ttk.LabelFrame(main_frame, text="Drop Files Here", padding="10")
        drop_frame.pack(fill=tk.X, pady=(0, 10))

        self.drop_label = ttk.Label(drop_frame,
                                   text="Drag and drop image files (PNG, JPG, BMP, TIFF) or PDF files here\n"
                                        "Or click 'Browse Files' to select files",
                                   justify=tk.CENTER,
                                   background="#3c3c3c",
                                   relief="sunken")
        self.drop_label.pack(fill=tk.X, expand=True, pady=20, padx=20)

        # Control buttons frame
        button_frame = ttk.Frame(main_frame)
        button_frame.pack(fill=tk.X, pady=(0, 10))

        self.browse_button = ttk.Button(button_frame, text="Browse Files",
                                       command=self.browse_files)
        self.browse_button.pack(side=tk.LEFT, padx=(0, 10))

        self.clear_button = ttk.Button(button_frame, text="Clear Results",
                                      command=self.clear_results, state=tk.DISABLED)
        self.clear_button.pack(side=tk.LEFT, padx=(0, 10))

        self.save_button = ttk.Button(button_frame, text="Save Results",
                                     command=self.save_results, state=tk.DISABLED)
        self.save_button.pack(side=tk.LEFT, padx=(0, 10))

        self.copy_button = ttk.Button(button_frame, text="Copy to Clipboard",
                                     command=self.copy_to_clipboard, state=tk.DISABLED)
        self.copy_button.pack(side=tk.LEFT, padx=(10, 0))

        self.clear_log_button = ttk.Button(button_frame, text="Clear Log",
                                          command=self.clear_log)
        self.clear_log_button.pack(side=tk.LEFT, padx=(10, 0))

        # Batch queue buttons
        batch_button_frame = ttk.Frame(main_frame)
        batch_button_frame.pack(fill=tk.X, pady=(0, 10))

        self.batch_start_button = ttk.Button(batch_button_frame, text="Start Batch",
                                            command=self.start_batch)
        self.batch_start_button.pack(side=tk.LEFT, padx=(0, 10))

        self.batch_cancel_button = ttk.Button(batch_button_frame, text="Cancel Batch",
                                             command=self.cancel_batch, state=tk.DISABLED)
        self.batch_cancel_button.pack(side=tk.LEFT, padx=(0, 10))

        self.batch_clear_button = ttk.Button(batch_button_frame, text="Clear Completed",
                                            command=self.clear_batch_completed)
        self.batch_clear_button.pack(side=tk.LEFT, padx=(0, 10))

        # Options frame with preprocessing and extraction toggles
        options_frame = ttk.LabelFrame(main_frame, text="Options", padding="10")
        options_frame.pack(fill=tk.X, pady=(0, 10))

        # Post-process checkbox
        self.post_process_var = tk.BooleanVar(value=self.config.get('post_process', False))
        self.post_process_checkbox = ttk.Checkbutton(
            options_frame, text="Post-process output",
            variable=self.post_process_var
        )
        self.post_process_checkbox.grid(row=0, column=0, sticky=tk.W, padx=(0, 20))

        # Extract tables checkbox
        self.extract_tables_var = tk.BooleanVar(value=self.config.get('extract_tables', False))
        self.extract_tables_checkbox = ttk.Checkbutton(
            options_frame, text="Extract tables from PDF",
            variable=self.extract_tables_var
        )
        self.extract_tables_checkbox.grid(row=0, column=1, sticky=tk.W, padx=(0, 20))

        # Preserve layout checkbox
        self.preserve_layout_var = tk.BooleanVar(value=self.config.get('preserve_layout', False))
        self.preserve_layout_checkbox = ttk.Checkbutton(
            options_frame, text="Preserve layout (DBSCAN)",
            variable=self.preserve_layout_var
        )
        self.preserve_layout_checkbox.grid(row=0, column=2, sticky=tk.W, padx=(0, 20))

        # Searchable PDF checkbox
        self.searchable_pdf_var = tk.BooleanVar(value=self.config.get('searchable_pdf', False))
        self.searchable_pdf_checkbox = ttk.Checkbutton(
            options_frame, text="Searchable PDF",
            variable=self.searchable_pdf_var
        )
        self.searchable_pdf_checkbox.grid(row=0, column=3, sticky=tk.W)

        # Preprocessing options row
        self.deskew_var = tk.BooleanVar(value=self.config.get('deskew', False))
        self.clahe_var = tk.BooleanVar(value=self.config.get('clahe', False))
        self.upscale_var = tk.BooleanVar(value=self.config.get('upscale', False))
        self.remove_bg_var = tk.BooleanVar(value=self.config.get('remove_bg', False))

        preprocess_frame = ttk.LabelFrame(main_frame, text="Preprocessing", padding="10")
        preprocess_frame.pack(fill=tk.X, pady=(0, 10))

        ttk.Checkbutton(preprocess_frame, text="Deskew",
                        variable=self.deskew_var).grid(row=0, column=0, sticky=tk.W, padx=(0, 20))
        ttk.Checkbutton(preprocess_frame, text="CLAHE",
                        variable=self.clahe_var).grid(row=0, column=1, sticky=tk.W, padx=(0, 20))
        ttk.Checkbutton(preprocess_frame, text="Upscale (2x)",
                        variable=self.upscale_var).grid(row=0, column=2, sticky=tk.W, padx=(0, 20))
        ttk.Checkbutton(preprocess_frame, text="Remove background",
                        variable=self.remove_bg_var).grid(row=0, column=3, sticky=tk.W)

        # Preview canvas
        preview_frame = ttk.LabelFrame(main_frame, text="Preview", padding="5")
        preview_frame.pack(fill=tk.X, pady=(0, 10))

        preview_inner = ttk.Frame(preview_frame)
        preview_inner.pack(fill=tk.X)

        self.preview_original_label = ttk.Label(preview_inner, text="Original image will appear here")
        self.preview_original_label.pack(side=tk.LEFT, padx=10)

        self.preview_processed_label = ttk.Label(preview_inner, text="Processed image will appear here")
        self.preview_processed_label.pack(side=tk.RIGHT, padx=10)

        # Progress bar
        self.progress_var = tk.DoubleVar()
        self.progress_bar = ttk.Progressbar(main_frame, variable=self.progress_var,
                                           maximum=100, mode='determinate')
        self.progress_bar.pack(fill=tk.X, pady=(0, 10))

        # Status bar with memory monitor
        status_bar = ttk.Frame(main_frame)
        status_bar.pack(fill=tk.X, pady=(0, 5))

        self.status_label = ttk.Label(status_bar, text="Ready")
        self.status_label.pack(side=tk.LEFT, anchor=tk.W)

        self.memory_label = ttk.Label(status_bar, text="Memory: N/A")
        self.memory_label.pack(side=tk.RIGHT, anchor=tk.E)

        # Notebook for results, log, and batch tabs
        self.notebook = ttk.Notebook(main_frame)
        self.notebook.pack(fill=tk.BOTH, expand=True)

        # Results frame (tab 1)
        results_frame = ttk.Frame(self.notebook)
        self.notebook.add(results_frame, text="Results")

        # Results text area
        self.results_text = scrolledtext.ScrolledText(results_frame, wrap=tk.WORD,
                                                      font=("Courier", 10),
                                                      bg='#1e1e1e', fg='#d4d4d4',
                                                      insertbackground='#d4d4d4')
        self.results_text.pack(fill=tk.BOTH, expand=True)

        # File info label
        self.file_info_label = ttk.Label(results_frame, text="")
        self.file_info_label.pack(anchor=tk.W, pady=(5, 0))

        # Log frame (tab 2)
        log_frame = ttk.Frame(self.notebook)
        self.notebook.add(log_frame, text="Log")

        self.log_text = scrolledtext.ScrolledText(log_frame, wrap=tk.WORD,
                                                    font=("Courier", 9),
                                                    state=tk.DISABLED, height=8,
                                                    bg='#1e1e1e', fg='#d4d4d4')
        self.log_text.pack(fill=tk.BOTH, expand=True)

        # Batch queue frame (tab 3)
        batch_frame = ttk.Frame(self.notebook)
        self.notebook.add(batch_frame, text="Batch Queue")

        # Batch queue treeview
        batch_tree_frame = ttk.Frame(batch_frame)
        batch_tree_frame.pack(fill=tk.BOTH, expand=True)

        self.batch_tree = ttk.Treeview(batch_tree_frame, columns=('status', 'progress'),
                                       show='headings', height=8)
        self.batch_tree.heading('#0', text='File')
        self.batch_tree.heading('status', text='Status')
        self.batch_tree.heading('progress', text='Progress')
        self.batch_tree.column('status', width=100)
        self.batch_tree.column('progress', width=100)

        batch_scroll = ttk.Scrollbar(batch_tree_frame, orient=tk.VERTICAL,
                                     command=self.batch_tree.yview)
        self.batch_tree.configure(yscrollcommand=batch_scroll.set)

        self.batch_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        batch_scroll.pack(side=tk.RIGHT, fill=tk.Y)

        # Batch results text
        self.batch_results_text = scrolledtext.ScrolledText(batch_frame, wrap=tk.WORD,
                                                           font=("Courier", 9),
                                                           height=8,
                                                           bg='#1e1e1e', fg='#d4d4d4')
        self.batch_results_text.pack(fill=tk.BOTH, expand=True)

    def setup_drag_drop(self):
        """Setup drag and drop functionality."""
        if not TKINTERDND_AVAILABLE:
            return

        self.drop_label.drop_target_register(DND_FILES)
        self.drop_label.dnd_bind('<<Drop>>', self.on_drop)
        self.drop_label.dnd_bind('<<DragEnter>>', self.on_drag_enter)
        self.drop_label.dnd_bind('<<DragLeave>>', self.on_drag_leave)

    def on_drag_enter(self, event):
        self.drop_label.config(background="#505050")

    def on_drag_leave(self, event):
        self.drop_label.config(background="#3c3c3c")

    def on_drop(self, event):
        """Handle file drop event."""
        try:
            files = self.root.splitlist(event.data)
            if files:
                file_path = Path(files[0])
                if file_path.exists():
                    self.process_file(file_path)
                else:
                    messagebox.showerror("File Error", f"Dropped file not found: {file_path}")
        except Exception as e:
            messagebox.showerror("Drop Error", f"Failed to process dropped file: {e}")

    def browse_files(self):
        """Open file browser dialog."""
        filetypes = [
            ('All supported files', '*.png *.jpg *.jpeg *.bmp *.tiff *.tif *.pdf'),
            ('Image files', '*.png *.jpg *.jpeg *.bmp *.tiff *.tif'),
            ('PDF files', '*.pdf'),
            ('All files', '*.*')
        ]

        filenames = filedialog.askopenfilenames(
            title="Select files to process",
            filetypes=filetypes
        )

        if filenames:
            if len(filenames) == 1:
                self.process_file(Path(filenames[0]))
            else:
                # Add multiple files to batch queue
                for filename in filenames:
                    self.add_to_batch(Path(filename))

    def _on_language_change(self, event=None):
        """Handle language selection change."""
        selected = self.language_var.get()
        lang_code = selected.split(' - ')[0] if ' - ' in selected else selected
        self.ocr_engine.languages = [lang_code]
        self._log(f"Language changed to: {lang_code}")

    def _on_confidence_change(self, value=None):
        """Handle confidence slider change."""
        threshold = self.confidence_var.get()
        self.confidence_label.config(text=f"{threshold:.2f}")
        self.ocr_engine.set_confidence_threshold(threshold)

    def add_to_batch(self, file_path: Path):
        """Add a file to the batch queue."""
        if not self.batch_manager:
            self._init_batch_manager()

        task_id = self.batch_manager.add_file(file_path)
        self.batch_tree.insert('', tk.END, iid=task_id, text=file_path.name,
                              values=('pending', '0%'))
        self._log(f"Added to batch queue: {file_path.name}")

    def _init_batch_manager(self):
        """Initialize the batch queue manager."""
        config = {
            'post_process': self.post_process_var.get(),
            'extract_tables': self.extract_tables_var.get(),
        }
        self.batch_manager = BatchQueueManager(
            ocr_engine=self.ocr_engine,
            pdf_handler=self.pdf_handler,
            config=config,
            on_progress=self._on_batch_progress,
            on_status_change=self._on_batch_status_change,
            on_task_complete=self._on_batch_task_complete,
            on_queue_complete=self._on_batch_queue_complete,
            on_log=self._log,
        )

    def start_batch(self):
        """Start batch processing."""
        if not self.batch_manager or self.batch_manager.get_pending_count() == 0:
            messagebox.showinfo("Batch Queue", "No files in batch queue. Use 'Browse Files' to add multiple files.")
            return

        self.batch_start_button.config(state=tk.DISABLED)
        self.batch_cancel_button.config(state=tk.NORMAL)
        self.batch_manager.start()

    def cancel_batch(self):
        """Cancel batch processing."""
        if self.batch_manager:
            self.batch_manager.cancel()
        self.batch_start_button.config(state=tk.NORMAL)
        self.batch_cancel_button.config(state=tk.DISABLED)

    def clear_batch_completed(self):
        """Clear completed batch tasks."""
        if self.batch_manager:
            self.batch_manager.clear_completed()
        # Clear treeview
        if hasattr(self, 'batch_tree'):
            for item in self.batch_tree.get_children():
                self.batch_tree.delete(item)
        self.batch_results_text.delete(1.0, tk.END)

    def _on_batch_progress(self, task_id: str, progress: float):
        """Handle batch task progress update."""
        def update():
            try:
                if self.batch_tree.exists(task_id):
                    self.batch_tree.item(task_id, values=('processing', f'{progress:.0f}%'))
            except tk.TclError:
                pass
        self.root.after(0, update)

    def _on_batch_status_change(self, task_id: str, status: str):
        """Handle batch task status change."""
        def update():
            try:
                if self.batch_tree.exists(task_id):
                    task = self.batch_manager.get_task(task_id) if self.batch_manager else None
                    progress = f'{task.progress:.0f}%' if task else ''
                    self.batch_tree.item(task_id, values=(status, progress))
            except tk.TclError:
                pass
        self.root.after(0, update)

    def _on_batch_task_complete(self, task_id: str, success: bool, result_text: str):
        """Handle batch task completion."""
        def update():
            try:
                if self.batch_tree.exists(task_id):
                    status = 'completed' if success else 'failed'
                    self.batch_tree.item(task_id, values=(status, '100%'))
                if result_text:
                    task = self.batch_manager.get_task(task_id) if self.batch_manager else None
                    filename = task.file_path.name if task else 'Unknown'
                    self.batch_results_text.insert(tk.END, f"=== {filename} ===\n{result_text}\n\n")
            except tk.TclError:
                pass
        self.root.after(0, update)

    def _on_batch_queue_complete(self, total: int, successful: int, failed: int):
        """Handle batch queue completion."""
        def update():
            self.batch_start_button.config(state=tk.NORMAL)
            self.batch_cancel_button.config(state=tk.DISABLED)
            messagebox.showinfo("Batch Complete",
                              f"Batch processing complete.\n"
                              f"Total: {total}, Successful: {successful}, Failed: {failed}")
        self.root.after(0, update)

    def process_file(self, file_path: Path):
        """Process a single file in a separate thread."""
        if self.processing:
            messagebox.showwarning("Processing", "Please wait for current processing to complete.")
            return

        try:
            if not file_path.exists():
                messagebox.showerror("File Error", f"File not found: {file_path}\n\nPlease check the file path and ensure the file exists.")
                return
        except OSError as e:
            messagebox.showerror("File Access Error", f"Cannot access file {file_path}: {e}\n\nThis may be due to permission issues or file system problems.")
            return

        # Check file extension
        supported_exts = {'.png', '.jpg', '.jpeg', '.bmp', '.tiff', '.tif', '.pdf'}
        if file_path.suffix.lower() not in supported_exts:
            messagebox.showerror("Unsupported File",
                                f"Unsupported file type: {file_path.suffix}\n\n"
                                "Supported formats: PNG, JPG, JPEG, BMP, TIFF, TIF, PDF\n\n"
                                "Please select a file with one of these extensions.")
            return

        # Start processing
        self.processing = True
        self.current_file = file_path
        self.status_label.config(text=f"Processing: {file_path.name}")
        self.progress_var.set(0)
        self.results_text.delete(1.0, tk.END)
        self.file_info_label.config(text="")
        self.clear_button.config(state=tk.DISABLED)
        self.save_button.config(state=tk.DISABLED)

        # Update preview
        self._update_preview(file_path)

        # Process in separate thread
        thread = threading.Thread(target=self._process_file_thread, args=(file_path,))
        thread.daemon = True
        thread.start()

    def _update_preview(self, file_path: Path):
        """Update the preview canvas with original and processed images."""
        try:
            from PIL import Image, ImageTk
            import cv2

            if file_path.suffix.lower() in {'.png', '.jpg', '.jpeg', '.bmp', '.tiff', '.tif'}:
                # Load original image
                img = Image.open(file_path)
                img.thumbnail((300, 200))
                self.preview_image_original = ImageTk.PhotoImage(img)
                self.preview_original_label.config(image=self.preview_image_original, text="")

                # Process image for preview
                img_cv = cv2.imread(str(file_path))
                if img_cv is not None:
                    gray = cv2.cvtColor(img_cv, cv2.COLOR_BGR2GRAY)
                    blurred = cv2.GaussianBlur(gray, (5, 5), 0)
                    thresh = cv2.adaptiveThreshold(
                        blurred, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C, cv2.THRESH_BINARY, 11, 2
                    )
                    img_rgb = cv2.cvtColor(thresh, cv2.COLOR_GRAY2RGB)
                    img_pil = Image.fromarray(img_rgb)
                    img_pil.thumbnail((300, 200))
                    self.preview_image_processed = ImageTk.PhotoImage(img_pil)
                    self.preview_processed_label.config(image=self.preview_image_processed, text="")
        except Exception as e:
            logger.debug(f"Preview update failed: {e}")

    def _process_file_thread(self, file_path: Path):
        """Process file in background thread with improved error handling."""
        try:
            # Update progress
            self.processing_queue.put(('progress', 10))
            self._log(f"Starting processing of {file_path.name}")

            # Determine engine
            engine = None
            if self.engine_var.get() != 'auto':
                engine = self.engine_var.get()

            # Update OCR engine settings from GUI
            self.ocr_engine.deskew = self.deskew_var.get()
            self.ocr_engine.clahe = self.clahe_var.get()
            self.ocr_engine.upscale = self.upscale_var.get()
            self.ocr_engine.remove_bg = self.remove_bg_var.get()
            self.ocr_engine.preserve_layout = self.preserve_layout_var.get()

            # Determine file type and process
            if file_path.suffix.lower() == '.pdf':
                self.processing_queue.put(('status', f"Detecting PDF type: {file_path.name}"))
                self._log("Detecting PDF type...")

                page_count = self.pdf_handler.get_page_count(file_path)
                pdf_type = self.pdf_handler.detect_pdf_type(file_path)
                self._log(f"PDF type: {pdf_type}, pages: {page_count}")
                self.processing_queue.put(('progress', 30))

                self.processing_queue.put(('status', f"Extracting text from {page_count} pages"))
                self._log("Extracting text from PDF...")

                text = self.pdf_handler.extract_text(file_path, engine=engine)
                self.processing_queue.put(('progress', 90))

                results = []
                metadata = {
                    'file_path': str(file_path),
                    'file_type': 'PDF',
                    'page_count': page_count,
                    'type': pdf_type,
                    'languages': self.ocr_engine.get_supported_languages(),
                    'confidence_threshold': self.ocr_engine.confidence_threshold,
                }

                # Table extraction
                if self.extract_tables_var.get():
                    self._log("Extracting tables from PDF...")
                    try:
                        tables = self.pdf_handler.extract_tables(file_path)
                        table_output = _format_gui_tables(tables)
                        if table_output:
                            text = text + "\n\n" + table_output
                    except Exception as e:
                        self._log(f"Table extraction failed: {e}")

                # Searchable PDF
                if self.searchable_pdf_var.get():
                    self._log("Generating searchable PDF...")
                    try:
                        search_path = self.pdf_handler.generate_searchable_pdf(
                            file_path, self.ocr_engine
                        )
                        self._log(f"Searchable PDF saved to: {search_path}")
                    except Exception as e:
                        self._log(f"Searchable PDF generation failed: {e}")
            else:
                self.processing_queue.put(('status', f"Preprocessing image: {file_path.name}"))
                self.processing_queue.put(('progress', 30))

                self.processing_queue.put(('status', "Extracting text with OCR"))
                self._log(f"Extracting text using engine: {engine or 'auto'}")
                text, results = self.ocr_engine.extract_text(file_path, engine=engine,
                                                            file_type='image', page_count=1)
                self.processing_queue.put(('progress', 90))

                metadata = {
                    'file_path': str(file_path),
                    'file_type': 'Image',
                    'languages': self.ocr_engine.get_supported_languages(),
                    'confidence_threshold': self.ocr_engine.confidence_threshold
                }

            # Apply post-processing if enabled
            if self.post_process_var.get():
                self.processing_queue.put(('status', "Post-processing text"))
                self._log("Applying post-processing...")
                text = self.ocr_engine.post_process_text(
                    text, self.ocr_engine.get_supported_languages()
                )

            self.processing_queue.put(('progress', 100))
            self.processing_queue.put(('status', "Processing complete"))
            self._log("Processing complete.")
            self.processing_queue.put(('results', text, results, metadata))

        except OCREngineNotAvailableError as e:
            error_msg = f"No OCR engines available: {e}"
            if e.details and 'suggestion' in e.details:
                error_msg += f"\n\nSuggestion: {e.details['suggestion']}"
            logger.error(error_msg)
            self._log(f"ERROR: {error_msg}")
            self.processing_queue.put(('error', error_msg))
        except OCRProcessingError as e:
            error_msg = f"OCR processing failed: {e}"
            if e.details and 'suggestion' in e.details:
                error_msg += f"\n\nSuggestion: {e.details['suggestion']}"
            logger.error(error_msg)
            self._log(f"ERROR: {error_msg}")
            self.processing_queue.put(('error', error_msg))
        except ImageProcessingError as e:
            error_msg = f"Image processing failed: {e}"
            if e.details and 'suggestion' in e.details:
                error_msg += f"\n\nSuggestion: {e.details['suggestion']}"
            logger.error(error_msg)
            self._log(f"ERROR: {error_msg}")
            self.processing_queue.put(('error', error_msg))
        except PDFProcessingError as e:
            error_msg = f"PDF processing failed: {e}"
            if e.details and 'suggestion' in e.details:
                error_msg += f"\n\nSuggestion: {e.details['suggestion']}"
            logger.error(error_msg)
            self._log(f"ERROR: {error_msg}")
            self.processing_queue.put(('error', error_msg))
        except MemoryError as e:
            error_msg = f"Memory limit exceeded: {e}"
            if e.details and 'suggestion' in e.details:
                error_msg += f"\n\nSuggestion: {e.details['suggestion']}"
            logger.error(error_msg)
            self._log(f"ERROR: {error_msg}")
            self.processing_queue.put(('error', error_msg))
        except FileNotFoundError as e:
            error_msg = f"File not found: {file_path}. Please check the file path and ensure it exists."
            logger.error(error_msg)
            self._log(f"ERROR: {error_msg}")
            self.processing_queue.put(('error', error_msg))
        except PermissionError as e:
            error_msg = f"Permission denied accessing {file_path}. Please check file permissions."
            logger.error(error_msg)
            self._log(f"ERROR: {error_msg}")
            self.processing_queue.put(('error', error_msg))
        except IsADirectoryError as e:
            error_msg = f"Path {file_path} is a directory, not a file. Please select a file."
            logger.error(error_msg)
            self._log(f"ERROR: {error_msg}")
            self.processing_queue.put(('error', error_msg))
        except OSError as e:
            error_msg = f"OS error processing {file_path}: {e}. This may be due to file system issues."
            logger.error(error_msg)
            self._log(f"ERROR: {error_msg}")
            self.processing_queue.put(('error', error_msg))
        except ValueError as e:
            error_msg = f"Invalid file format or corrupted file: {file_path}. Error: {e}"
            logger.error(error_msg)
            self._log(f"ERROR: {error_msg}")
            self.processing_queue.put(('error', error_msg))
        except RuntimeError as e:
            error_msg = f"OCR engine error: {e}"
            logger.error(error_msg)
            self._log(f"ERROR: {error_msg}")
            self.processing_queue.put(('error', error_msg))
        except Exception as e:
            error_msg = f"Unexpected error processing {file_path}: {e}"
            logger.error(error_msg)
            self._log(f"ERROR: {error_msg}")
            self.processing_queue.put(('error', error_msg))

    def process_queue(self):
        """Process messages from the background thread."""
        try:
            while True:
                message = self.processing_queue.get_nowait()
                msg_type = message[0]

                if msg_type == 'progress':
                    self.progress_var.set(message[1])
                elif msg_type == 'status':
                    self.status_label.config(text=message[1])
                elif msg_type == 'results':
                    text, results, metadata = message[1], message[2], message[3]
                    self.display_results(text, results, metadata)
                elif msg_type == 'error':
                    error_msg = message[1]
                    self.status_label.config(text=f"Error: {error_msg}")
                    messagebox.showerror("Processing Error", error_msg)
                    self.reset_ui()
                elif msg_type == 'log':
                    log_entry = message[1]
                    self.log_text.config(state=tk.NORMAL)
                    self.log_text.insert(tk.END, log_entry + '\n')
                    self.log_text.see(tk.END)
                    self.log_text.config(state=tk.DISABLED)

        except queue.Empty:
            pass

        # Schedule next check
        self.root.after(100, self.process_queue)

    def display_results(self, text: str, results: List[Dict], metadata: Dict[str, Any]):
        """Display processing results in the UI."""
        # Display text
        self.results_text.delete(1.0, tk.END)
        self.results_text.insert(tk.END, text)

        # Display file info
        info_parts = []
        if metadata.get('file_type') == 'PDF':
            info_parts.append(f"Pages: {metadata.get('page_count', 'Unknown')}")
            info_parts.append(f"Type: {metadata.get('type', 'Unknown')}")
        else:
            info_parts.append(f"Languages: {', '.join(metadata.get('languages', ['en']))}")
            info_parts.append(f"Confidence: {metadata.get('confidence_threshold', 0.5)}")

        if self.post_process_var.get():
            info_parts.append("Post-processed")

        self.file_info_label.config(text=" | ".join(info_parts))

        # Update UI state
        self.clear_button.config(state=tk.NORMAL)
        self.save_button.config(state=tk.NORMAL)
        self.copy_button.config(state=tk.NORMAL)
        self.processing = False
        self.status_label.config(text="Ready")

    def clear_results(self):
        """Clear the results display."""
        self.results_text.delete(1.0, tk.END)
        self.file_info_label.config(text="")
        self.clear_button.config(state=tk.DISABLED)
        self.save_button.config(state=tk.DISABLED)
        self.copy_button.config(state=tk.DISABLED)
        self.progress_var.set(0)
        self.status_label.config(text="Ready")

    def clear_log(self):
        """Clear the log display."""
        self.log_text.config(state=tk.NORMAL)
        self.log_text.delete(1.0, tk.END)
        self.log_text.config(state=tk.DISABLED)

    def save_results(self):
        """Save the current results to a file."""
        text = self.results_text.get(1.0, tk.END).strip()
        if not text:
            messagebox.showwarning("Save Results", "No results to save.")
            return

        filetypes = [
            ('Text files', '*.txt'),
            ('Markdown files', '*.md'),
            ('JSON files', '*.json'),
            ('CSV files', '*.csv'),
            ('HTML files', '*.html'),
            ('Word documents', '*.docx'),
            ('All files', '*.*')
        ]

        filename = filedialog.asksaveasfilename(
            title="Save Results",
            defaultextension=".txt",
            filetypes=filetypes
        )

        if filename:
            try:
                suffix = Path(filename).suffix.lower()
                if suffix == '.docx':
                    from docx import Document
                    doc = Document()
                    if self.current_file:
                        doc.add_heading('OCR Results', level=1)
                        doc.add_paragraph(f"Original file: {self.current_file}")
                        doc.add_paragraph(f"Languages: {', '.join(self.ocr_engine.get_supported_languages())}")
                    doc.add_paragraph(text)
                    doc.save(filename)
                else:
                    with open(filename, 'w', encoding='utf-8') as f:
                        if self.current_file:
                            f.write(f"# OCR Results\n")
                            f.write(f"# Original file: {self.current_file}\n")
                            f.write(f"# Languages: {', '.join(self.ocr_engine.get_supported_languages())}\n")
                            f.write(f"# Processed with OCR Text Extractor v{__import__('src').__version__}\n\n")
                        f.write(text)
                messagebox.showinfo("Save Results", f"Results saved to {filename}")
            except ImportError:
                messagebox.showerror("Save Error", "python-docx is required to save as .docx. Install with: pip install python-docx")
            except PermissionError as e:
                messagebox.showerror("Save Error", f"Permission denied saving to {filename}.\n\nPlease check directory permissions and try a different location.")
            except OSError as e:
                messagebox.showerror("Save Error", f"OS error saving to {filename}: {e}\n\nThis may be due to disk space or file system issues.")
            except Exception as e:
                messagebox.showerror("Save Error", f"Unexpected error saving to {filename}: {e}")

    def copy_to_clipboard(self):
        """Copy current results to clipboard."""
        text = self.results_text.get(1.0, tk.END).strip()
        if not text:
            messagebox.showwarning("Copy to Clipboard", "No results to copy.")
            return

        try:
            self.root.clipboard_clear()
            self.root.clipboard_append(text)
            self.status_label.config(text="Copied to clipboard")
            messagebox.showinfo("Copy to Clipboard", "Results copied to clipboard.")
        except Exception as e:
            messagebox.showerror("Copy Error", f"Failed to copy to clipboard: {e}")

    def reset_ui(self):
        """Reset UI to ready state after error."""
        self.processing = False
        self.progress_var.set(0)
        self.status_label.config(text="Ready")


def main():
    """Main entry point for the GUI application."""
    if not TKINTERDND_AVAILABLE:
        print("WARNING: tkinterdnd2 not available. Drag-and-drop functionality will be disabled.")
        print("Install with: pip install tkinterdnd2")
        root = tk.Tk()
    else:
        root = TkinterDnD.Tk()

    app = OCRGUIApp(root)
    root.mainloop()


if __name__ == '__main__':
    main()
