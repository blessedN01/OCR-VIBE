#!/usr/bin/env python3
"""
Flet GUI Application Module

Modern desktop GUI for OCR-VIBE using Flet, replacing the legacy tkinter GUI.
"""

import logging
import os
import queue
import tempfile
import threading
import time
import datetime
import asyncio
from pathlib import Path
from typing import Any, Dict, List, Optional

import psutil

import flet as ft
from .ocr_engine import OCREngine, get_available_engines
from .pdf_handler import PDFHandler
from .config_loader import get_config_parser
from .batch_queue import BatchQueueManager
from .exceptions import (
    OCREngineNotAvailableError,
    OCRProcessingError,
    ImageProcessingError,
    PDFProcessingError,
    MemoryError,
)
from .cli_app import OutputFormatter, parse_page_ranges

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

AVAILABLE_LANGUAGES = [
    ('en', 'English'),
    ('fr', 'French'),
    ('de', 'German'),
    ('es', 'Spanish'),
    ('it', 'Italian'),
    ('pt', 'Portuguese'),
    ('nl', 'Dutch'),
    ('ru', 'Russian'),
    ('ja', 'Japanese'),
    ('ko', 'Korean'),
    ('zh', 'Chinese'),
    ('ar', 'Arabic'),
    ('hi', 'Hindi'),
    ('tr', 'Turkish'),
    ('pl', 'Polish'),
    ('cs', 'Czech'),
    ('sv', 'Swedish'),
    ('da', 'Danish'),
    ('fi', 'Finnish'),
    ('el', 'Greek'),
]


class FletGUIApp:
    def __init__(self, page: ft.Page):
        self.page = page
        self.page.title = "OCR-VIBE"
        self.page.theme_mode = ft.ThemeMode.DARK
        self.page.window_width = 1100
        self.page.window_height = 850
        self.page.window_min_width = 900
        self.page.window_min_height = 700
        self.page.padding = 20

        self.config = get_config_parser()
        self.ocr_engine = OCREngine(
            languages=self.config.get('languages', ['en']),
            confidence_threshold=self.config.get('confidence_threshold', 0.5),
            use_gpu=self.config.get('use_gpu', False),
            preprocess_images=self.config.get('preprocess_images', True),
            memory_limit_mb=self.config.get('memory_limit_mb'),
            deskew=self.config.get('deskew', False),
            clahe=self.config.get('clahe', False),
            upscale=self.config.get('upscale', False),
            remove_bg=self.config.get('remove_bg', False),
            preserve_layout=self.config.get('preserve_layout', False),
            auto_select=self.config.get('auto_select', True),
        )
        self.pdf_handler = PDFHandler(ocr_engine=self.ocr_engine)

        self.processing = False
        self.current_file: Optional[Path] = None
        self.log_messages: List[str] = []
        self.batch_manager: Optional[BatchQueueManager] = None
        self.batch_rows: Dict[str, ft.Row] = {}
        self._temp_files: List[str] = []

        self.queue: queue.Queue = queue.Queue()
        self._ui_dirty = False

        self._build_ui()
        self.page.run_thread(self._drain_queue_loop)
        self.page.run_thread(self._update_memory_usage_loop)

    def _build_ui(self):
        self.title_text = ft.Text("OCR-VIBE", size=24, weight=ft.FontWeight.BOLD)

        self.language_dropdown = ft.Dropdown(
            label="Language",
            value=self.config.get('languages', ['en'])[0] if self.config.get('languages') else 'en',
            options=[ft.dropdown.Option(code, name) for code, name in AVAILABLE_LANGUAGES],
            width=220,
        )
        self.confidence_slider = ft.Slider(
            min=0.0,
            max=1.0,
            value=self.config.get('confidence_threshold', 0.5),
            label="{value:.2f}",
            width=220,
        )
        engines = get_available_engines()
        engine_options = [ft.Radio(value="auto", label="Auto")] + [
            ft.Radio(value=e, label=e.capitalize()) for e in engines
        ]
        self.engine_radio_group = ft.RadioGroup(content=ft.Row(engine_options), value="auto")

        self.output_format_dropdown = ft.Dropdown(
            label="Output format",
            value=self.config.get('output_format', 'txt'),
            options=[
                ft.dropdown.Option("txt", "TXT"),
                ft.dropdown.Option("json", "JSON"),
                ft.dropdown.Option("csv", "CSV"),
                ft.dropdown.Option("md", "Markdown"),
                ft.dropdown.Option("html", "HTML"),
                ft.dropdown.Option("docx", "DOCX"),
            ],
            width=140,
        )

        self.page_range_field = ft.TextField(
            label="PDF pages (e.g. 1-3,5,7-9)",
            value="",
            width=200,
            hint_text="All pages if empty",
        )

        ocr_config_tab = ft.Column(
            [
                ft.Text("OCR Options", size=16, weight=ft.FontWeight.BOLD),
                ft.Row(
                    [
                        ft.Column([self.language_dropdown], tight=True),
                        ft.Column([ft.Text("Confidence"), self.confidence_slider], tight=True),
                        ft.Column([ft.Text("Engine"), self.engine_radio_group], tight=True),
                        ft.Column([self.output_format_dropdown], tight=True),
                        ft.Column([self.page_range_field], tight=True),
                    ],
                    alignment=ft.MainAxisAlignment.START,
                    spacing=20,
                    wrap=True,
                ),
            ],
            spacing=10,
        )

        self.drop_container = ft.Container(
            content=ft.Text(
                "Drop files here or click Browse\n"
                "Supports PNG, JPG, BMP, TIFF, PDF",
                text_align=ft.TextAlign.CENTER,
                size=14,
            ),
            border=ft.Border.all(2, ft.Colors.with_opacity(0.5, ft.Colors.ON_SURFACE)),
            border_radius=8,
            padding=12,
            ink=True,
            on_click=self._browse_files,
            on_hover=self._on_hover_drop_zone,
        )

        self.browse_button = ft.ElevatedButton("Browse", on_click=self._browse_files)
        self.clear_button = ft.ElevatedButton("Clear Results", on_click=self._clear_results, disabled=True)
        self.save_button = ft.ElevatedButton("Save Results", on_click=self._save_results, disabled=True)
        self.copy_button = ft.ElevatedButton("Copy to Clipboard", on_click=self._copy_to_clipboard, disabled=True)
        self.clear_log_button = ft.ElevatedButton("Clear Log", on_click=self._clear_log)

        action_row = ft.Row(
            [self.browse_button, self.clear_button, self.save_button, self.copy_button, self.clear_log_button],
            spacing=10,
            wrap=True,
        )

        self.folder_button = ft.ElevatedButton("Browse Folder", on_click=self._browse_folder)
        self.recursive_checkbox = ft.Checkbox(label="Recursive", value=False)
        self.folder_row = ft.Row(
            [self.folder_button, self.recursive_checkbox],
            spacing=10,
            visible=False,
        )

        self.gpu_checkbox = ft.Checkbox(label="Use GPU", value=self.config.get('use_gpu', False))
        self.memory_limit_field = ft.TextField(
            label="Memory limit (MB)",
            value=str(self.config.get('memory_limit_mb', '')),
            width=140,
        )
        self.preprocess_global_checkbox = ft.Checkbox(label="Enable preprocessing", value=self.config.get('preprocess_images', True))

        system_tab = ft.Column(
            [
                ft.Text("System", size=16, weight=ft.FontWeight.BOLD),
                ft.Row(
                    [self.gpu_checkbox, self.memory_limit_field, self.preprocess_global_checkbox],
                    spacing=20,
                ),
            ],
            spacing=10,
        )

        self.batch_start_button = ft.ElevatedButton("Start Batch", on_click=self._start_batch)
        self.batch_cancel_button = ft.ElevatedButton("Cancel Batch", on_click=self._cancel_batch, disabled=True)
        self.batch_clear_button = ft.ElevatedButton("Clear Completed", on_click=self._clear_batch_completed)
        self.batch_parallel_checkbox = ft.Checkbox(label="Parallel processing", value=False)

        batch_row = ft.Row(
            [self.batch_start_button, self.batch_cancel_button, self.batch_clear_button, self.batch_parallel_checkbox],
            spacing=10,
            wrap=True,
        )

        self.post_process_checkbox = ft.Checkbox(label="Post-process output", value=self.config.get('post_process', False))
        self.extract_tables_checkbox = ft.Checkbox(label="Extract tables from PDF", value=self.config.get('extract_tables', False))
        self.preserve_layout_checkbox = ft.Checkbox(label="Preserve layout (DBSCAN)", value=self.config.get('preserve_layout', False))
        self.searchable_pdf_checkbox = ft.Checkbox(label="Searchable PDF", value=self.config.get('searchable_pdf', False))

        pdf_options_tab = ft.Column(
            [
                ft.Text("PDF Options", size=16, weight=ft.FontWeight.BOLD),
                ft.Row(
                    [self.post_process_checkbox, self.extract_tables_checkbox, self.preserve_layout_checkbox, self.searchable_pdf_checkbox],
                    spacing=20,
                    wrap=True,
                ),
            ],
            spacing=10,
        )

        self.deskew_checkbox = ft.Checkbox(label="Deskew", value=self.config.get('deskew', False))
        self.clahe_checkbox = ft.Checkbox(label="CLAHE", value=self.config.get('clahe', False))
        self.upscale_checkbox = ft.Checkbox(label="Upscale (2x)", value=self.config.get('upscale', False))
        self.remove_bg_checkbox = ft.Checkbox(label="Remove background", value=self.config.get('remove_bg', False))

        preprocessing_tab = ft.Column(
            [
                ft.Text("Preprocessing", size=16, weight=ft.FontWeight.BOLD),
                ft.Row(
                    [self.deskew_checkbox, self.clahe_checkbox, self.upscale_checkbox, self.remove_bg_checkbox],
                    spacing=20,
                ),
            ],
            spacing=10,
        )

        self.preview_original_image = ft.Image(src="", width=300, height=200, fit=ft.BoxFit.CONTAIN, visible=False)
        self.preview_processed_image = ft.Image(src="", width=300, height=200, fit=ft.BoxFit.CONTAIN, visible=False)
        self.preview_original_text = ft.Text("Original image will appear here", size=12, color=ft.Colors.ON_SURFACE_VARIANT)
        self.preview_processed_text = ft.Text("Processed image will appear here", size=12, color=ft.Colors.ON_SURFACE_VARIANT)

        preview_row = ft.Row(
            [
                ft.Column([self.preview_original_text, self.preview_original_image]),
                ft.Column([self.preview_processed_text, self.preview_processed_image]),
            ],
            spacing=20,
            visible=False,
        )
        self.preview_row = preview_row

        preview_tab = ft.Column(
            [
                ft.Text("Preview", size=16, weight=ft.FontWeight.BOLD),
                preview_row,
            ],
            spacing=10,
        )

        self.progress_bar = ft.ProgressBar(value=0.0, width=float("inf"))

        self.status_text = ft.Text("Ready", size=12)
        self.memory_text = ft.Text("Memory: N/A", size=12)
        status_bar = ft.Row([self.status_text, ft.Container(expand=True), self.memory_text], alignment=ft.MainAxisAlignment.SPACE_BETWEEN)

        self.results_textfield = ft.TextField(
            multiline=True,
            read_only=True,
            border=ft.InputBorder.NONE,
            bgcolor=ft.Colors.with_opacity(0.0, ft.Colors.SURFACE),
            text_style=ft.TextStyle(font_family="Courier", size=12),
            expand=True,
        )

        self.log_textfield = ft.TextField(
            multiline=True,
            read_only=True,
            border=ft.InputBorder.NONE,
            bgcolor=ft.Colors.with_opacity(0.0, ft.Colors.SURFACE),
            text_style=ft.TextStyle(font_family="Courier", size=11),
            expand=True,
        )

        self.batch_column = ft.Column([], spacing=5, expand=True)
        self.batch_results_textfield = ft.TextField(
            multiline=True,
            read_only=True,
            border=ft.InputBorder.NONE,
            bgcolor=ft.Colors.with_opacity(0.0, ft.Colors.SURFACE),
            text_style=ft.TextStyle(font_family="Courier", size=11),
            expand=True,
        )

        results_tab = ft.Column(
            [ft.Container(self.results_textfield, padding=10, expand=True)],
            expand=True,
            scroll=ft.ScrollMode.AUTO,
        )
        log_tab = ft.Column(
            [ft.Container(self.log_textfield, padding=10, expand=True)],
            expand=True,
            scroll=ft.ScrollMode.AUTO,
        )
        batch_tab = ft.Column(
            [
                ft.Text("Batch Queue", size=16, weight=ft.FontWeight.BOLD),
                self.batch_column,
                self.batch_results_textfield,
            ],
            expand=True,
            spacing=10,
            scroll=ft.ScrollMode.AUTO,
        )

        self.tabs = ft.Tabs(
            selected_index=0,
            animation_duration=300,
            expand=True,
            length=9,
            content=ft.Column(
                expand=True,
                controls=[
                    ft.TabBar(
                        tabs=[
                            ft.Tab(label="Results"),
                            ft.Tab(label="Log"),
                            ft.Tab(label="Batch Queue"),
                            ft.Tab(label="OCR Options"),
                            ft.Tab(label="PDF Options"),
                            ft.Tab(label="Preprocessing"),
                            ft.Tab(label="Preview"),
                            ft.Tab(label="System"),
                            ft.Tab(label="About"),
                        ],
                    ),
                    ft.TabBarView(
                        expand=True,
                        controls=[
                            results_tab,
                            log_tab,
                            batch_tab,
                            ocr_config_tab,
                            pdf_options_tab,
                            preprocessing_tab,
                            preview_tab,
                            system_tab,
                            self._build_about_tab(),
                        ],
                    ),
                ],
            ),
        )

        self.file_picker = ft.FilePicker()
        self.page.services.append(self.file_picker)

        self.save_file_picker = ft.FilePicker()
        self.page.services.append(self.save_file_picker)

        self.folder_picker = ft.FilePicker()
        self.page.services.append(self.folder_picker)

        top_section = ft.Column(
            [
                self.title_text,
                self.drop_container,
                action_row,
                self.folder_row,
                batch_row,
                self.progress_bar,
                status_bar,
            ],
            spacing=6,
        )

        self.column = ft.Column(
            [
                top_section,
                ft.Container(self.tabs, expand=True),
            ],
            spacing=0,
            expand=True,
        )

    def _on_hover_drop_zone(self, e: ft.HoverEvent):
        if e.data == "true":
            self.drop_container.border = ft.Border.all(2, ft.Colors.PRIMARY)
            self.drop_container.bgcolor = ft.Colors.with_opacity(0.1, ft.Colors.PRIMARY)
        else:
            self.drop_container.border = ft.Border.all(2, ft.Colors.with_opacity(0.5, ft.Colors.ON_SURFACE))
            self.drop_container.bgcolor = ft.Colors.with_opacity(0.0, ft.Colors.SURFACE)
        self.page.update()

    async def _browse_files(self, e: ft.ControlEvent):
        try:
            files = await self.file_picker.pick_files(
                dialog_title="Select files to process",
                file_type=ft.FilePickerFileType.CUSTOM,
                allowed_extensions=["png", "jpg", "jpeg", "bmp", "tiff", "tif", "pdf"],
            )
            if files:
                paths = [Path(f.path) for f in files]
                if len(paths) == 1:
                    self._process_file(paths[0])
                else:
                    for fp in paths:
                        self._add_to_batch(fp)
        except Exception as ex:
            logger.error(f"File picker error: {ex}")
            self._show_snack(f"Error selecting files: {ex}")

    async def _browse_folder(self, e: ft.ControlEvent):
        try:
            result = await self.folder_picker.get_directory_path(
                dialog_title="Select folder to process",
            )
            if result:
                folder_path = Path(result)
                recursive = self.recursive_checkbox.value
                self._process_folder(folder_path, recursive)
        except Exception as ex:
            logger.error(f"Folder picker error: {ex}")
            self._show_snack(f"Error selecting folder: {ex}")

    def _process_folder(self, folder_path: Path, recursive: bool):
        if self.processing:
            self._show_snack("Please wait for current processing to complete.")
            return
        supported_extensions = {'.png', '.jpg', '.jpeg', '.bmp', '.tiff', '.tif', '.pdf'}
        if recursive:
            files = [f for f in folder_path.rglob('*') if f.is_file() and f.suffix.lower() in supported_extensions]
        else:
            files = [f for f in folder_path.iterdir() if f.is_file() and f.suffix.lower() in supported_extensions]

        if not files:
            self._show_snack("No supported files found in folder.")
            return

        self.processing = True
        self.status_text.value = f"Processing folder: {folder_path.name} ({len(files)} files)"
        self.progress_bar.value = 0.0
        self.results_textfield.value = ""
        self.clear_button.disabled = True
        self.save_button.disabled = True
        self.copy_button.disabled = True
        self.preview_row.visible = False
        self.page.update()

        thread = threading.Thread(target=self._process_folder_thread, args=(folder_path, files), daemon=True)
        thread.start()

    def _process_file(self, file_path: Path):
        if self.processing:
            self._show_snack("Please wait for current processing to complete.")
            return
        if not file_path.exists():
            self._show_snack(f"File not found: {file_path}")
            return
        if file_path.suffix.lower() not in {'.png', '.jpg', '.jpeg', '.bmp', '.tiff', '.tif', '.pdf'}:
            self._show_snack(f"Unsupported file type: {file_path.suffix}")
            return

        self.processing = True
        self.current_file = file_path
        self.status_text.value = f"Processing: {file_path.name}"
        self.progress_bar.value = 0.0
        self.results_textfield.value = ""
        self.clear_button.disabled = True
        self.save_button.disabled = True
        self.copy_button.disabled = True
        self.preview_row.visible = False
        self.preview_original_image.visible = False
        self.preview_processed_image.visible = False
        self.preview_original_text.value = "Original image will appear here"
        self.preview_processed_text.value = "Processed image will appear here"
        self.page.update()

        self._update_preview(file_path)

        thread = threading.Thread(target=self._process_file_thread, args=(file_path,), daemon=True)
        thread.start()

    def _update_preview(self, file_path: Path):
        try:
            if file_path.suffix.lower() in {'.png', '.jpg', '.jpeg', '.bmp', '.tiff', '.tif'}:
                import cv2
                from PIL import Image

                img = Image.open(file_path)
                img.thumbnail((300, 200))
                self.preview_original_image.src = str(file_path)
                self.preview_original_image.visible = True
                self.preview_original_text.value = "Original"
                self.preview_original_text.visible = True

                img_cv = cv2.imread(str(file_path))
                if img_cv is not None:
                    gray = cv2.cvtColor(img_cv, cv2.COLOR_BGR2GRAY)
                    blurred = cv2.GaussianBlur(gray, (5, 5), 0)
                    thresh = cv2.adaptiveThreshold(
                        blurred, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C, cv2.THRESH_BINARY, 11, 2
                    )
                    img_rgb = cv2.cvtColor(thresh, cv2.COLOR_GRAY2RGB)
                    img_pil = Image.fromarray(img_rgb)
                    img_pil.thumbnail((300, 200))
                    tmp = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
                    img_pil.save(tmp.name)
                    self.preview_processed_image.src = tmp.name
                    self.preview_processed_image.visible = True
                    self.preview_processed_text.value = "Processed"
                    self.preview_processed_text.visible = True
                    self.preview_processed_image.data = tmp.name
                    self._temp_files.append(tmp.name)

                self.preview_row.visible = True
                self.page.update()
        except Exception as ex:
            logger.debug(f"Preview update failed: {ex}")

    def _process_file_thread(self, file_path: Path):
        try:
            self.queue.put(("progress", 10))
            self._log(f"Starting processing of {file_path.name}")

            engine = None
            if self.engine_radio_group.value != "auto":
                engine = self.engine_radio_group.value

            self.ocr_engine.deskew = self.deskew_checkbox.value
            self.ocr_engine.clahe = self.clahe_checkbox.value
            self.ocr_engine.upscale = self.upscale_checkbox.value
            self.ocr_engine.remove_bg = self.remove_bg_checkbox.value
            self.ocr_engine.preserve_layout = self.preserve_layout_checkbox.value
            self.ocr_engine.set_confidence_threshold(self.confidence_slider.value)

            page_range = self.page_range_field.value.strip() if self.page_range_field.value else None
            page_numbers = None
            if page_range:
                try:
                    page_numbers = parse_page_ranges(page_range)
                except ValueError as ex:
                    self.queue.put(("error", f"Invalid page range: {ex}"))
                    self._reset_ui()
                    return

            if file_path.suffix.lower() == '.pdf':
                self.queue.put(("status", f"Detecting PDF type: {file_path.name}"))
                self._log("Detecting PDF type...")

                page_count = self.pdf_handler.get_page_count(file_path)
                pdf_type = self.pdf_handler.detect_pdf_type(file_path)
                self._log(f"PDF type: {pdf_type}, pages: {page_count}")
                self.queue.put(("progress", 30))

                self.queue.put(("status", f"Extracting text from {page_count} pages"))
                self._log("Extracting text from PDF...")

                text = self.pdf_handler.extract_text(file_path, engine=engine, pages=page_numbers)
                self.queue.put(("progress", 90))

                results = []
                metadata = {
                    'file_path': str(file_path),
                    'file_type': 'PDF',
                    'page_count': page_count,
                    'type': pdf_type,
                    'languages': self.ocr_engine.get_supported_languages(),
                    'confidence_threshold': self.ocr_engine.confidence_threshold,
                }

                if self.extract_tables_checkbox.value:
                    self._log("Extracting tables from PDF...")
                    try:
                        tables = self.pdf_handler.extract_tables(file_path, pages=','.join(str(p) for p in page_numbers) if page_numbers else None)
                        table_output = self._format_gui_tables(tables)
                        if table_output:
                            text = text + "\n\n" + table_output
                    except Exception as ex:
                        self._log(f"Table extraction failed: {ex}")

                if self.searchable_pdf_checkbox.value:
                    self._log("Generating searchable PDF...")
                    try:
                        search_path = self.pdf_handler.generate_searchable_pdf(file_path, self.ocr_engine, pages=page_numbers)
                        self._log(f"Searchable PDF saved to: {search_path}")
                    except Exception as ex:
                        self._log(f"Searchable PDF generation failed: {ex}")
            else:
                self.queue.put(("status", f"Preprocessing image: {file_path.name}"))
                self.queue.put(("progress", 30))

                self.queue.put(("status", "Extracting text with OCR"))
                self._log(f"Extracting text using engine: {engine or 'auto'}")
                text, results = self.ocr_engine.extract_text(
                    file_path, engine=engine, file_type='image', page_count=1
                )
                self.queue.put(("progress", 90))

                metadata = {
                    'file_path': str(file_path),
                    'file_type': 'Image',
                    'languages': self.ocr_engine.get_supported_languages(),
                    'confidence_threshold': self.ocr_engine.confidence_threshold,
                }

            if self.post_process_checkbox.value:
                self.queue.put(("status", "Post-processing text"))
                self._log("Applying post-processing...")
                text = self.ocr_engine.post_process_text(text, self.ocr_engine.get_supported_languages())

            self.queue.put(("progress", 100))
            self.queue.put(("status", "Processing complete"))
            self._log("Processing complete.")
            self.queue.put(("results", text, results, metadata))

        except OCREngineNotAvailableError as ex:
            error_msg = f"No OCR engines available: {ex}"
            if ex.details and 'suggestion' in ex.details:
                error_msg += f"\n\nSuggestion: {ex.details['suggestion']}"
            logger.error(error_msg)
            self._log(f"ERROR: {error_msg}")
            self.queue.put(("error", error_msg))
        except OCRProcessingError as ex:
            error_msg = f"OCR processing failed: {ex}"
            if ex.details and 'suggestion' in ex.details:
                error_msg += f"\n\nSuggestion: {ex.details['suggestion']}"
            logger.error(error_msg)
            self._log(f"ERROR: {error_msg}")
            self.queue.put(("error", error_msg))
        except ImageProcessingError as ex:
            error_msg = f"Image processing failed: {ex}"
            if ex.details and 'suggestion' in ex.details:
                error_msg += f"\n\nSuggestion: {ex.details['suggestion']}"
            logger.error(error_msg)
            self._log(f"ERROR: {error_msg}")
            self.queue.put(("error", error_msg))
        except PDFProcessingError as ex:
            error_msg = f"PDF processing failed: {ex}"
            if ex.details and 'suggestion' in ex.details:
                error_msg += f"\n\nSuggestion: {ex.details['suggestion']}"
            logger.error(error_msg)
            self._log(f"ERROR: {error_msg}")
            self.queue.put(("error", error_msg))
        except MemoryError as ex:
            error_msg = f"Memory limit exceeded: {ex}"
            if ex.details and 'suggestion' in ex.details:
                error_msg += f"\n\nSuggestion: {ex.details['suggestion']}"
            logger.error(error_msg)
            self._log(f"ERROR: {error_msg}")
            self.queue.put(("error", error_msg))
        except FileNotFoundError as ex:
            error_msg = f"File not found: {file_path}. Please check the file path and ensure it exists."
            logger.error(error_msg)
            self._log(f"ERROR: {error_msg}")
            self.queue.put(("error", error_msg))
        except PermissionError as ex:
            error_msg = f"Permission denied accessing {file_path}. Please check file permissions."
            logger.error(error_msg)
            self._log(f"ERROR: {error_msg}")
            self.queue.put(("error", error_msg))
        except IsADirectoryError as ex:
            error_msg = f"Path {file_path} is a directory, not a file. Please select a file."
            logger.error(error_msg)
            self._log(f"ERROR: {error_msg}")
            self.queue.put(("error", error_msg))
        except OSError as ex:
            error_msg = f"OS error processing {file_path}: {ex}. This may be due to file system issues."
            logger.error(error_msg)
            self._log(f"ERROR: {error_msg}")
            self.queue.put(("error", error_msg))
        except ValueError as ex:
            error_msg = f"Invalid file format or corrupted file: {file_path}. Error: {ex}"
            logger.error(error_msg)
            self._log(f"ERROR: {error_msg}")
            self.queue.put(("error", error_msg))
        except RuntimeError as ex:
            error_msg = f"OCR engine error: {ex}"
            logger.error(error_msg)
            self._log(f"ERROR: {error_msg}")
            self.queue.put(("error", error_msg))
        except Exception as ex:
            error_msg = f"Unexpected error processing {file_path}: {ex}"
            logger.error(error_msg)
            self._log(f"ERROR: {error_msg}")
            self.queue.put(("error", error_msg))

    def _process_folder_thread(self, folder_path: Path, files: list[Path]):
        try:
            output_format = self.output_format_dropdown.value
            page_range = self.page_range_field.value.strip() if self.page_range_field.value else None
            page_numbers = None
            if page_range:
                try:
                    page_numbers = parse_page_ranges(page_range)
                except ValueError as ex:
                    self.queue.put(("error", f"Invalid page range: {ex}"))
                    self._reset_ui()
                    return

            from .cli_app import process_single_file

            ocr_engine = self.ocr_engine
            pdf_handler = self.pdf_handler

            all_results = []
            successful = 0
            total = len(files)

            for i, file_path in enumerate(files):
                if not self.processing:
                    break
                progress = (i / total) * 100.0
                self.queue.put(("progress", progress))
                self.queue.put(("status", f"Processing {file_path.name} ({i+1}/{total})"))
                self._log(f"Processing {file_path.name}...")

                try:
                    text, _results = process_single_file(
                        file_path=file_path,
                        ocr_engine=ocr_engine,
                        pdf_handler=pdf_handler,
                        output_format=output_format,
                        output_path=None,
                        engine=self.engine_radio_group.value if self.engine_radio_group.value != "auto" else None,
                        post_process=self.post_process_checkbox.value,
                        extract_tables=self.extract_tables_checkbox.value,
                        preserve_layout=self.preserve_layout_checkbox.value,
                        searchable_pdf=self.searchable_pdf_checkbox.value,
                        pages=page_numbers,
                    )
                    successful += 1
                    all_results.append(f"=== {file_path.name} ===\n{text}\n")
                except Exception as ex:
                    self._log(f"Failed {file_path.name}: {ex}")

            self.queue.put(("progress", 100))
            self.queue.put(("status", f"Folder complete: {successful}/{total} successful"))
            combined = "\n\n".join(all_results)
            self.queue.put(("results", combined, [], {
                'file_path': str(folder_path),
                'file_type': 'Folder',
                'languages': self.ocr_engine.get_supported_languages(),
                'confidence_threshold': self.ocr_engine.confidence_threshold,
                'output_format': output_format,
            }))
        except Exception as ex:
            error_msg = f"Folder processing error: {ex}"
            logger.error(error_msg)
            self._log(f"ERROR: {error_msg}")
            self.queue.put(("error", error_msg))

    def _drain_queue(self):
        try:
            while True:
                message = self.queue.get_nowait()
                msg_type = message[0]
                self._ui_dirty = True

                if msg_type == "progress":
                    self.progress_bar.value = message[1] / 100.0
                elif msg_type == "status":
                    self.status_text.value = message[1]
                elif msg_type == "memory":
                    self.memory_text.value = message[1]
                elif msg_type == "results":
                    text, results, metadata = message[1], message[2], message[3]
                    self._display_results(text, results, metadata)
                elif msg_type == "error":
                    error_msg = message[1]
                    self.status_text.value = f"Error: {error_msg}"
                    self._show_snack(error_msg)
                    self._reset_ui()
                elif msg_type == "log":
                    log_entry = message[1]
                    if self.log_textfield.value:
                        self.log_textfield.value += log_entry + "\n"
                    else:
                        self.log_textfield.value = log_entry + "\n"
                elif msg_type == "batch_progress":
                    task_id, progress = message[1], message[2]
                    row = self.batch_rows.get(task_id)
                    if row:
                        row.controls[2].value = progress / 100.0
                elif msg_type == "batch_status":
                    task_id, status = message[1], message[2]
                    row = self.batch_rows.get(task_id)
                    if row:
                        row.controls[1].value = status
                elif msg_type == "batch_complete":
                    task_id, success, result_text = message[1], message[2], message[3]
                    row = self.batch_rows.get(task_id)
                    if row:
                        row.controls[1].value = "completed" if success else "failed"
                        row.controls[2].value = 1.0
                    if result_text:
                        if self.batch_results_textfield.value:
                            self.batch_results_textfield.value += "\n\n" + result_text
                        else:
                            self.batch_results_textfield.value = result_text
                elif msg_type == "batch_queue_complete":
                    self.batch_start_button.disabled = False
                    self.batch_cancel_button.disabled = True
                    total, successful, failed = message[1], message[2], message[3]
                    self._show_snack(f"Batch complete: {successful}/{total} successful, {failed} failed")
                elif msg_type == "batch_log":
                    if self.log_textfield.value:
                        self.log_textfield.value += message[1] + "\n"
                    else:
                        self.log_textfield.value = message[1] + "\n"
        except queue.Empty:
            pass

        if self._ui_dirty:
            self.page.update()
            self._ui_dirty = False

    def _drain_queue_loop(self):
        while True:
            self._drain_queue()
            time.sleep(0.1)

    def _update_memory_usage_loop(self):
        while True:
            self._update_memory_usage()
            time.sleep(1.0)

    def _display_results(self, text: str, results: List[dict], metadata: Dict[str, Any]):
        self.results_textfield.value = text
        info_parts = []
        if metadata.get('file_type') == 'PDF':
            info_parts.append(f"Pages: {metadata.get('page_count', 'Unknown')}")
            info_parts.append(f"Type: {metadata.get('type', 'Unknown')}")
        else:
            info_parts.append(f"Languages: {', '.join(metadata.get('languages', ['en']))}")
            info_parts.append(f"Confidence: {metadata.get('confidence_threshold', 0.5)}")

        if self.post_process_checkbox.value:
            info_parts.append("Post-processed")

        self.status_text.value = " | ".join(info_parts)
        self.clear_button.disabled = False
        self.save_button.disabled = False
        self.copy_button.disabled = False
        self.processing = False
        self.progress_bar.value = 1.0
        self.page.update()

    def _clear_results(self, e: ft.ControlEvent):
        self._cleanup_temp_files()
        self.results_textfield.value = ""
        self.status_text.value = "Ready"
        self.clear_button.disabled = True
        self.save_button.disabled = True
        self.copy_button.disabled = True
        self.progress_bar.value = 0.0
        self.page.update()

    def _cleanup_temp_files(self):
        for tmp_path in self._temp_files:
            try:
                os.unlink(tmp_path)
            except OSError:
                pass
        self._temp_files.clear()

    def _clear_log(self, e: ft.ControlEvent):
        self.log_textfield.value = ""
        self.page.update()

    async def _save_results(self, e: ft.ControlEvent):
        text = self.results_textfield.value.strip()
        if not text:
            self._show_snack("No results to save.")
            return
        path = await self.save_file_picker.save_file(
            dialog_title="Save Results",
            file_name="ocr_results.txt",
            file_type=ft.FilePickerFileType.CUSTOM,
            allowed_extensions=["txt", "md", "json", "csv", "html", "docx"],
        )
        if not path:
            return
        try:
            suffix = Path(path).suffix.lower()
            output_format = self.output_format_dropdown.value
            if suffix == ".docx":
                try:
                    from docx import Document
                except ImportError:
                    self._show_snack("python-docx is required to save as .docx")
                    return
                doc = Document()
                if self.current_file:
                    doc.add_heading("OCR Results", level=1)
                    doc.add_paragraph(f"Original file: {self.current_file}")
                    doc.add_paragraph(f"Languages: {', '.join(self.ocr_engine.get_supported_languages())}")
                doc.add_paragraph(text)
                doc.save(path)
            else:
                metadata = {
                    'file_path': str(self.current_file) if self.current_file else 'Unknown',
                    'languages': self.ocr_engine.get_supported_languages(),
                    'confidence_threshold': self.ocr_engine.confidence_threshold,
                    'output_format': output_format,
                }
                if output_format == 'json':
                    content = OutputFormatter.format_json(text, [], metadata)
                elif output_format == 'csv':
                    content = OutputFormatter.format_csv(text, [], metadata)
                elif output_format == 'md':
                    content = OutputFormatter.format_md(text, metadata)
                elif output_format == 'html':
                    content = OutputFormatter.format_html(text, [], metadata)
                else:
                    content = OutputFormatter.format_txt(text, metadata)
                with open(path, "w", encoding="utf-8") as f:
                    f.write(content)
            self._show_snack(f"Results saved to {path}")
        except Exception as ex:
            self._show_snack(f"Save error: {ex}")

    def _copy_to_clipboard(self, e: ft.ControlEvent):
        text = self.results_textfield.value.strip()
        if not text:
            self._show_snack("No results to copy.")
            return
        self.page.set_clipboard(text)
        self._show_snack("Results copied to clipboard.")

    def _add_to_batch(self, file_path: Path):
        if not self.batch_manager:
            self._init_batch_manager()

        task_id = self.batch_manager.add_file(file_path)
        row = ft.Row(
            [
                ft.Text(file_path.name, expand=True, size=12),
                ft.Text("pending", size=12, width=100),
                ft.ProgressBar(value=0.0, width=120),
            ],
            spacing=10,
        )
        self.batch_rows[task_id] = row
        self.batch_column.controls.append(row)
        self._log(f"Added to batch queue: {file_path.name}")
        self.page.update()

    def _init_batch_manager(self):
        memory_limit = None
        if self.memory_limit_field.value.strip():
            try:
                memory_limit = int(self.memory_limit_field.value.strip())
            except ValueError:
                pass

        config = {
            'post_process': self.post_process_checkbox.value,
            'extract_tables': self.extract_tables_checkbox.value,
            'engine': self.engine_radio_group.value if self.engine_radio_group.value != "auto" else None,
            'preserve_layout': self.preserve_layout_checkbox.value,
            'searchable_pdf': self.searchable_pdf_checkbox.value,
            'deskew': self.deskew_checkbox.value,
            'clahe': self.clahe_checkbox.value,
            'upscale': self.upscale_checkbox.value,
            'remove_bg': self.remove_bg_checkbox.value,
            'preprocess_images': self.preprocess_global_checkbox.value,
            'use_gpu': self.gpu_checkbox.value,
            'memory_limit_mb': memory_limit,
            'languages': self.ocr_engine.languages,
            'confidence_threshold': self.confidence_slider.value,
            'output_format': self.output_format_dropdown.value,
            'parallel': self.batch_parallel_checkbox.value,
        }
        self.batch_manager = BatchQueueManager(
            ocr_engine=self.ocr_engine,
            pdf_handler=self.pdf_handler,
            config=config,
            on_progress=self._on_batch_progress,
            on_status_change=self._on_batch_status_change,
            on_task_complete=self._on_batch_task_complete,
            on_queue_complete=self._on_batch_queue_complete,
            on_log=self._log,
        )

    def _start_batch(self, e: ft.ControlEvent):
        if not self.batch_manager or self.batch_manager.get_pending_count() == 0:
            self._show_snack("No files in batch queue. Use 'Browse' to add multiple files.")
            return
        self.batch_start_button.disabled = True
        self.batch_cancel_button.disabled = False
        self.batch_manager.start()

    def _cancel_batch(self, e: ft.ControlEvent):
        if self.batch_manager:
            self.batch_manager.cancel()
        self.batch_start_button.disabled = False
        self.batch_cancel_button.disabled = True
        self.page.update()

    def _clear_batch_completed(self, e: ft.ControlEvent):
        if self.batch_manager:
            cleared_ids = self.batch_manager.clear_completed()
            for tid in cleared_ids:
                row = self.batch_rows.pop(tid, None)
                if row and row in self.batch_column.controls:
                    self.batch_column.controls.remove(row)
        self.batch_results_textfield.value = ""
        self.page.update()

    def _on_batch_progress(self, task_id: str, progress: float):
        self.queue.put(("batch_progress", task_id, progress))

    def _on_batch_status_change(self, task_id: str, status: str):
        self.queue.put(("batch_status", task_id, status))

    def _on_batch_task_complete(self, task_id: str, success: bool, result_text: str):
        self.queue.put(("batch_complete", task_id, success, result_text))

    def _on_batch_queue_complete(self, total: int, successful: int, failed: int):
        self.queue.put(("batch_queue_complete", total, successful, failed))

    def _log(self, message: str):
        timestamp = datetime.datetime.now().strftime('%H:%M:%S')
        entry = f"[{timestamp}] {message}"
        self.log_messages.append(entry)
        if len(self.log_messages) > 1000:
            self.log_messages = self.log_messages[-500:]
        self.queue.put(("log", entry))

    def _update_memory_usage(self):
        try:
            process = psutil.Process(os.getpid())
            mem_mb = process.memory_info().rss / (1024 * 1024)
            mem_percent = process.memory_percent()
            self.queue.put(("memory", f"Memory: {mem_mb:.0f} MB ({mem_percent:.1f}%)"))
        except Exception:
            self.queue.put(("memory", "Memory: N/A"))

    def _reset_ui(self):
        self.processing = False
        self.progress_bar.value = 0.0
        self.status_text.value = "Ready"
        self.clear_button.disabled = True
        self.save_button.disabled = True
        self.copy_button.disabled = True
        self.page.update()

    def _show_snack(self, message: str):
        self.page.show_dialog(ft.SnackBar(ft.Text(message)))

    @staticmethod
    def _format_gui_tables(tables: list) -> str:
        if not tables:
            return ""
        output_parts = []
        for i, table in enumerate(tables):
            try:
                if not hasattr(table, 'df') or table.df is None:
                    continue
                text = table.df.to_string(index=False)
                output_parts.append(f"--- Table {i + 1} ---\n{text}")
            except Exception as ex:
                logger.warning(f"Failed to format table {i + 1}: {ex}")
        return "\n\n".join(output_parts)

    def _build_about_tab(self):
        engines = get_available_engines()
        engine_rows = []
        for e in engines:
            initialized = (
                (e == 'easyocr' and self.ocr_engine.easyocr_reader is not None)
                or (e == 'paddleocr' and self.ocr_engine.paddleocr_reader is not None)
            )
            if not initialized and e == 'tesseract':
                from .ocr_engine import TESSERACT_AVAILABLE
                initialized = TESSERACT_AVAILABLE
            status = "Ready" if initialized else "Not initialized"
            color = ft.Colors.GREEN if initialized else ft.Colors.ORANGE
            engine_rows.append(
                ft.Row([
                    ft.Text(e, width=120),
                    ft.Text(status, color=color, width=120),
                ], spacing=10)
            )

        about_content = ft.Column([
            ft.Text("About OCR-VIBE", size=18, weight=ft.FontWeight.BOLD),
            ft.Text("Version 2.0.0", size=14),
            ft.Divider(),
            ft.Text("Available OCR Engines", size=16, weight=ft.FontWeight.BOLD),
            ft.Column(engine_rows, spacing=5),
            ft.Divider(),
            ft.Text("Features", size=16, weight=ft.FontWeight.BOLD),
            ft.Text("• Multi-format: PNG, JPG, JPEG, TIFF, BMP, PDF"),
            ft.Text("• Multiple OCR engines: EasyOCR, PaddleOCR, Tesseract"),
            ft.Text("• Batch and folder processing"),
            ft.Text("• Export: TXT, JSON, CSV, Markdown, HTML, DOCX"),
            ft.Text("• PDF: table extraction, searchable PDF, page range"),
            ft.Text("• Preprocessing: deskew, CLAHE, upscale, background removal"),
            ft.Text("• Layout preservation, language post-processing"),
            ft.Divider(),
            ft.Text("Configuration: ocr-vibe.yaml / ~/.ocr-vibe/config.yaml"),
        ], spacing=10, scroll=ft.ScrollMode.AUTO)
        return ft.Container(about_content, padding=20, expand=True)

    def build(self) -> ft.Column:
        return self.column
